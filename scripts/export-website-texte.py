"""Exportiert alle Website-Texte in eine Excel-Tabelle (3 Spalten)."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROWS: list[tuple[str, str]] = [
    # Meta / Browser-Titel
    ("layout.tsx → metadata.title", "Frank Vullhorst – KI‑Sparring"),
    (
        "layout.tsx → metadata.description",
        "KI-Sparring auf Augenhöhe: In 45 Minuten klären wir, wo KI und Automatisierung in Ihrem Betrieb entlasten können – sicher, sinnvoll und strategisch.",
    ),
    ("layout.tsx → openGraph.title", "Frank Vullhorst – KI‑Sparring"),
    (
        "layout.tsx → openGraph.description",
        "Praxisnah, nachvollziehbar und an EU AI Act sowie DSGVO orientiert.",
    ),
    (
        "termin/page.tsx → metadata.title",
        "Erstgespräch – Klarheit in 45 Minuten | Frank Vullhorst",
    ),
    (
        "termin/page.tsx → metadata.description",
        "Kostenloses Erstgespräch: Wo KI in Ihrem Betrieb entlastet – und welcher nächste Schritt passt.",
    ),
    # Header
    ("site-chrome.tsx → Header → Logo", "Frank Vullhorst"),
    ("site-chrome.tsx → Nav → Link 1", "Für wen"),
    ("site-chrome.tsx → Nav → Link 2", "Angebot"),
    ("site-chrome.tsx → Nav → Link 3", "Methode"),
    ("site-chrome.tsx → Nav → Link 4", "FAQ"),
    ("site-chrome.tsx → Nav → Link 5", "Termin"),
    # Hero
    ("page.tsx → Hero → kicker", "Meine Einladung:"),
    ("page.tsx → Hero → Headline Zeile 1", "Ein Erstgespräch."),
    ("page.tsx → Hero → Headline Zeile 2", "45 Minuten Klarheit."),
    ("copy.ts → cta.primary.hero", "Meine Klarheit in 45 Minuten sichern"),
    (
        "page.tsx → Hero → Microcopy Zeile 1",
        "Sie erzählen, was gerade ist – ich höre zu und ordne ein.",
    ),
    (
        "page.tsx → Hero → Microcopy Zeile 2",
        "45 Minuten. Kein Sales, kein Druck. Kostenfrei. Unverbindlich.",
    ),
    (
        "copy.ts → triggers.heroAfter (nicht sichtbar)",
        "Kostenloses Erstgespräch · Keine Vorbereitung · Kein Verkaufsdruck · 30+ Jahre Praxis",
    ),
    (
        "copy.ts → cta.secondary (nicht sichtbar)",
        "Zuerst sehen, welcher Weg zu mir passt",
    ),
    # Über mich
    ("page.tsx → Über mich → Bild alt", "Frank Vullhorst"),
    ("page.tsx → Über mich → kicker", "Über mich:"),
    ("page.tsx → Über mich → Titel", "Mein Name ist Frank."),
    (
        "page.tsx → Über mich → Lead",
        "Ich arbeite mit Inhabern und Geschäftsführern in Handwerk und KMU, die KI nicht als Hype, sondern als Entlastung im Alltag verstehen wollen – verständlich, sicher und praxisnah.",
    ),
    (
        "page.tsx → Über mich → Fließtext",
        "KI sollte nicht zusätzlich belasten, sondern im Alltag wirklich helfen. Ich unterstütze Betriebe dabei, KI sicher und nachvollziehbar in ihre Abläufe zu integrieren – mit 30+ Jahren Praxis aus Technik, Prozessen und Umsetzung.",
    ),
    ("page.tsx → Über mich → Tagline 1", "KI"),
    ("page.tsx → Über mich → Tagline 2", "Sicher"),
    ("page.tsx → Über mich → Tagline 3", "Sinnvoll"),
    ("page.tsx → Über mich → Tagline 4", "Strategisch"),
    # Für wen
    ("page.tsx → Für wen → kicker", "Für wen:"),
    (
        "page.tsx → Für wen → Titel Zeile 1",
        "Für Entscheider in Handwerk und KMU,",
    ),
    (
        "page.tsx → Für wen → Titel Zeile 2",
        "die Klarheit statt Buzzwords wollen.",
    ),
    (
        "page.tsx → Für wen → Lead",
        "Inhaber, Geschäftsführer und Verantwortliche, die enorm viel halten – organisieren, entscheiden, umsetzen – und dabei KI pragmatisch einordnen wollen.",
    ),
    (
        "page.tsx → Für wen → Absatz 1",
        "Menschen, die wissen: KI kann entlasten – aber nur, wenn sie zu Prozessen, Daten und Verantwortung passt. Die keine Tool-Show brauchen, sondern einen Fahrplan, den man im Betrieb wirklich gehen kann.",
    ),
    (
        "page.tsx → Für wen → Absatz 2",
        "Wenn Sie zu wenig Zeit im Tagesgeschäft haben, unsicher sind, was sinnvoll ist, Schatten-KI im Team bemerken oder Datenschutz und EU AI Act pragmatisch lösen wollen – dann sind Sie hier richtig.",
    ),
    # Angebot
    ("page.tsx → Angebot → kicker", "Mein Angebot:"),
    ("page.tsx → Angebot → Titel Zeile 1", "Ihr Einstieg."),
    ("page.tsx → Angebot → Titel Zeile 2", "Schritt für Schritt."),
    (
        "page.tsx → Angebot → Lead",
        "Kein Leistungskatalog zum Durchklicken – sondern Bausteine, die bei Ihnen wirklich Nutzen bringen.",
    ),
    ("page.tsx → Angebot → Karte 1 Titel", "Erstgespräch"),
    (
        "page.tsx → Angebot → Karte 1 Beschreibung",
        "Ihr Thema auf den Tisch bringen, sehen, was wirklich da ist.",
    ),
    ("page.tsx → Angebot → Karte 1 Meta", "45 Minuten · kostenfrei"),
    ("page.tsx → Angebot → Karte 2 Titel", "Potenzial-Scan"),
    (
        "page.tsx → Angebot → Karte 2 Beschreibung",
        "Wo KI heute entlastet – und welcher nächste Schritt wirtschaftlich sinnvoll ist.",
    ),
    (
        "page.tsx → Angebot → Karte 2 Meta",
        "Individuell · online oder vor Ort",
    ),
    ("page.tsx → Angebot → Karte 3 Titel", "AI-ISCA Audit"),
    (
        "page.tsx → Angebot → Karte 3 Beschreibung",
        "Nachvollziehbare Bestandsaufnahme: KI-Nutzung, Datenschutz, Compliance, Kompetenz.",
    ),
    (
        "page.tsx → Angebot → Karte 3 Meta",
        "Strukturiert · als Entscheidungsgrundlage",
    ),
    (
        "page.tsx → Angebot → Karte 4 Titel",
        "KI-Leitlinie & Sicherer Hafen",
    ),
    (
        "page.tsx → Angebot → Karte 4 Beschreibung",
        "Klare Regeln fürs Team – und die passende Schutzstufe für Ihre Daten.",
    ),
    ("page.tsx → Angebot → Karte 4 Meta", "Praxisnah · ohne Overkill"),
    ("page.tsx → Angebot → Karte 5 Titel", "AI Literacy Workshop"),
    (
        "page.tsx → Angebot → Karte 5 Beschreibung",
        "Kompetenz im Team – verständlich, nachweisbar, rechtssicher (Art. 4 EU AI Act).",
    ),
    ("page.tsx → Angebot → Karte 5 Meta", "Module 1–4 · für Ihr Team"),
    ("page.tsx → Angebot → Karte 6 Titel", "Pilotprojekt"),
    (
        "page.tsx → Angebot → Karte 6 Beschreibung",
        "Erster messbarer Nutzen im Betrieb – kontrolliert und auswertbar.",
    ),
    ("page.tsx → Angebot → Karte 6 Meta", "Klein starten · sauber ausbauen"),
    ("copy.ts → cta.offerInline (Button alle Karten)", "Termin wählen"),
    (
        "copy.ts → cta.primary.angebot (nicht sichtbar)",
        "Diese Klarheit für meinen Betrieb holen",
    ),
    ("copy.ts → triggers.angebotTitle", "Was Sie nach 45 Minuten haben"),
    (
        "copy.ts → triggers.angebotBody",
        "Eine konkrete Einordnung für Ihren Betrieb – und eine Empfehlung, die wirklich zu Ihnen passt.",
    ),
    (
        "copy.ts → triggers.angebotProof",
        "Cert-IT KI-Manager · Rhein-Main · KMU & Handwerk",
    ),
    # Methode
    ("page.tsx → Methode → kicker", "Wie ich arbeite:"),
    (
        "page.tsx → Methode → Titel Zeile 1",
        "Ich eröffne Reflexionsräume.",
    ),
    (
        "page.tsx → Methode → Titel Zeile 2",
        "Praxisnah und strukturiert.",
    ),
    (
        "page.tsx → Methode → Lead",
        "Fünf Schritte – nicht als Theorie, sondern als Fahrplan, den Sie in Ihrem Tempo und mit Ihrem Team gehen können.",
    ),
    ("page.tsx → Methode → Schritt 01 Nummer", "01"),
    ("page.tsx → Methode → Schritt 01 Titel", "Audit"),
    (
        "page.tsx → Methode → Schritt 01 Text",
        "Sie sehen klar: Was läuft, was fehlt, wo Risiko steckt.",
    ),
    ("page.tsx → Methode → Schritt 02 Nummer", "02"),
    ("page.tsx → Methode → Schritt 02 Titel", "KI-Leitlinie"),
    (
        "page.tsx → Methode → Schritt 02 Text",
        "Ihr Team weiß: Was darf genutzt werden – und was nicht.",
    ),
    ("page.tsx → Methode → Schritt 03 Nummer", "03"),
    ("page.tsx → Methode → Schritt 03 Titel", "Sicherer Hafen"),
    (
        "page.tsx → Methode → Schritt 03 Text",
        "Daten bekommen die Schutzstufe, die sie brauchen – ohne Overkill.",
    ),
    ("page.tsx → Methode → Schritt 04 Nummer", "04"),
    ("page.tsx → Methode → Schritt 04 Titel", "AI Literacy"),
    (
        "page.tsx → Methode → Schritt 04 Text",
        "Kompetenz im Team – nachweisbar, rechtssicher (Art. 4 EU AI Act).",
    ),
    ("page.tsx → Methode → Schritt 05 Nummer", "05"),
    ("page.tsx → Methode → Schritt 05 Titel", "Pilotprojekt"),
    (
        "page.tsx → Methode → Schritt 05 Text",
        "Erster messbarer Nutzen im Betrieb – kontrolliert und auswertbar.",
    ),
    # Versprechen
    ("page.tsx → Versprechen → kicker", "Das Versprechen:"),
    ("page.tsx → Versprechen → Titel Zeile 1", "Vom Unsichersein"),
    ("page.tsx → Versprechen → Titel Zeile 2", "zur Klarheit."),
    (
        "page.tsx → Versprechen → Lead",
        "Klarheit entsteht, wenn wir ehrlich hinsehen. Sie werden Worte finden für das, was Sie innerlich schon ahnen – und einen nächsten Schritt, der zu Ihrem Betrieb passt.",
    ),
    (
        "page.tsx → Versprechen → Absatz 1",
        "Damit Sie klar erkennen: wo KI heute bereits sinnvoll entlastet, welche Lösungen wirklich zu Ihrem Unternehmen passen – und welcher nächste Schritt wirtschaftlich sinnvoll ist.",
    ),
    (
        "page.tsx → Versprechen → Absatz 2 (kursiv)",
        "Nach unserem Gespräch haben Sie keine Buzzwords mehr – sondern Orientierung, konkrete Möglichkeiten und eine realistische Einschätzung für Ihr Unternehmen.",
    ),
    # Vertrauen
    ("page.tsx → Vertrauen → kicker", "Vertrauen:"),
    ("page.tsx → Vertrauen → Titel Zeile 1", "Warum Entscheider"),
    ("page.tsx → Vertrauen → Titel Zeile 2", "mir vertrauen."),
    (
        "page.tsx → Vertrauen → Lead",
        "Keine Show, keine Buzzwords – sondern jemand, der Betrieb von innen kennt und KI so erklärt, dass Sie handeln können.",
    ),
    (
        "page.tsx → Vertrauen → Bullet 1",
        "30+ Jahre Praxis – Technik, Prozesse, Umsetzung in der Industrie",
    ),
    (
        "page.tsx → Vertrauen → Bullet 2",
        "KI-Manager Certificate (Cert-IT), Nr. KI001220 · 02/2026",
    ),
    (
        "page.tsx → Vertrauen → Bullet 3",
        "27 Jahre bei 3D Systems – internationale Projekte bis €1,5M Budget",
    ),
    (
        "page.tsx → Vertrauen → Bullet 4",
        "Fokus: verständlich, sicher, umsetzbar (EU AI Act & DSGVO)",
    ),
    (
        "page.tsx → Vertrauen → Zitat 1 Text",
        "»Frank bringt Struktur und Klarheit in komplexe Themen – und schafft einen Raum, in dem man offen reflektieren kann, ohne bewertet zu werden.«",
    ),
    (
        "page.tsx → Vertrauen → Zitat 1 Quelle",
        "Führungskraft, Produktionsbetrieb · KI-Einstieg",
    ),
    (
        "page.tsx → Vertrauen → Zitat 2 Text",
        "»Endlich jemand, der KI nicht als Hype verkauft, sondern pragmatisch einordnet – mit Blick auf Datenschutz und Alltag im Betrieb.«",
    ),
    (
        "page.tsx → Vertrauen → Zitat 2 Quelle",
        "Inhaber, Handwerksbetrieb · Potenzial-Scan",
    ),
    ("page.tsx → Vertrauen → Erstgespräch kicker", "Erstgespräch"),
    (
        "page.tsx → Vertrauen → Erstgespräch Titel",
        "Was Sie im Erstgespräch bekommen",
    ),
    ("page.tsx → Vertrauen → Punkt 1 Titel", "Klarheit"),
    (
        "page.tsx → Vertrauen → Punkt 1 Text",
        "Worum es bei Ihnen wirklich geht – ohne Nebel.",
    ),
    ("page.tsx → Vertrauen → Punkt 2 Titel", "Einordnung"),
    (
        "page.tsx → Vertrauen → Punkt 2 Text",
        "Was sinnvoll ist – und was Sie (noch) lassen können.",
    ),
    ("page.tsx → Vertrauen → Punkt 3 Titel", "Nächster Schritt"),
    (
        "page.tsx → Vertrauen → Punkt 3 Text",
        "Eine saubere Empfehlung für den kleinsten sinnvollen Einstieg.",
    ),
    (
        "copy.ts → cta.primary.vertrauen",
        "Mein kostenloses Erstgespräch wählen",
    ),
    (
        "copy.ts → triggers.vertrauenAfter",
        "So starten viele Entscheider: erst Klarheit, dann der passende Baustein.",
    ),
    # FAQ
    ("page.tsx → FAQ → kicker", "Fragen & Antworten:"),
    ("page.tsx → FAQ → Titel", "FAQ"),
    (
        "page.tsx → FAQ → Lead",
        "Kurz, verständlich, entscheidungsfreundlich.",
    ),
    ("page.tsx → FAQ → Frage 1", "Was passiert im Erstgespräch?"),
    (
        "page.tsx → FAQ → Antwort 1",
        "Wir klären Ihre Ausgangslage, ordnen Chancen und Risiken ein und definieren den nächsten sinnvollen Schritt – ohne Tool‑Show und ohne Druck.",
    ),
    ("page.tsx → FAQ → Frage 2", "Muss ich etwas vorbereiten?"),
    (
        "page.tsx → FAQ → Antwort 2",
        "Nein. Hilfreich ist nur ein grober Überblick: Welche Aufgaben kosten Zeit? Wo wird schon KI genutzt? Welche Daten sind sensibel?",
    ),
    ("page.tsx → FAQ → Frage 3", "Was ist AI‑ISCA?"),
    (
        "page.tsx → FAQ → Antwort 3",
        "Ein strukturiertes Assessment (Audit): KI‑Nutzung, Datenschutz, Compliance und Kompetenz im Betrieb – als nachvollziehbare Grundlage für Entscheidungen.",
    ),
    (
        "page.tsx → FAQ → Frage 4",
        "Wie ist das mit Datenschutz und EU AI Act?",
    ),
    (
        "page.tsx → FAQ → Antwort 4",
        "Wir arbeiten pragmatisch: wenig Daten, klare Regeln, passende Schutzstufe. Fachbegriffe erkläre ich kurz und verständlich.",
    ),
    (
        "page.tsx → FAQ → Frage 5",
        "Für wen ist das Angebot nicht geeignet?",
    ),
    (
        "page.tsx → FAQ → Antwort 5",
        "Wenn Sie nur eine Tool‑Demo oder „Hype‑Beratung“ suchen. Es geht um umsetzbare Entlastung im Alltag – nicht um Buzzwords.",
    ),
    (
        "page.tsx → FAQ → Frage 6",
        "Was ist der nächste Schritt nach dem Gespräch?",
    ),
    (
        "page.tsx → FAQ → Antwort 6",
        "Typisch sind ein Potenzial‑Scan, ein AI‑ISCA Audit oder ein passender Workshop‑Einstieg (Modul 1). Immer klein starten, sauber ausbauen.",
    ),
    # Final CTA
    ("page.tsx → Final CTA → kicker", "Gemeinsam starten:"),
    ("page.tsx → Final CTA → Titel Zeile 1", "Statt Unsicherheit:"),
    (
        "page.tsx → Final CTA → Titel Zeile 2",
        "ein klarer nächster Schritt.",
    ),
    (
        "page.tsx → Final CTA → Lead",
        "In 45 Minuten wissen Sie, was bei Ihnen sinnvoll ist – sicher, nachvollziehbar, umsetzbar. Ohne Hype, ohne Verpflichtung.",
    ),
    (
        "copy.ts → triggers.finalBefore",
        "Kostenlos und unverbindlich – Sie wählen einen Termin, der in Ihren Kalender passt.",
    ),
    (
        "copy.ts → cta.primary.final",
        "Meinen nächsten sinnvollen Schritt klären",
    ),
    (
        "copy.ts → cta.primary.footer (nicht sichtbar)",
        "Klarheit für meinen Betrieb holen",
    ),
    (
        "copy.ts → triggers.footerBody (nicht sichtbar)",
        "Viele Inhaber und Geschäftsführer starten mit einem kurzen Gespräch – und wissen danach, ob und wie es weitergeht.",
    ),
    # Footer
    (
        "site-chrome.tsx → Footer → kicker",
        "Ich bin da, in Roßdorf und darüber hinaus:",
    ),
    (
        "site-chrome.tsx → Footer → Headline",
        "Rufen Sie mich an oder schreiben Sie mir direkt.",
    ),
    ("site-chrome.tsx → Footer → Telefon", "+49 (0)172 6689960"),
    (
        "site-chrome.tsx → Footer → E-Mail",
        "Info@FrankVullhorst.de",
    ),
    ("site-chrome.tsx → Footer → Link Impressum", "Impressum"),
    ("site-chrome.tsx → Footer → Link Datenschutz", "Datenschutz"),
    (
        "site-chrome.tsx → Footer → Copyright",
        "© {aktuelles Jahr} Frank Vullhorst · KI-Sparring",
    ),
    # Termin
    ("termin/page.tsx → Zurück-Link", "← Zurück zur Startseite"),
    ("termin/page.tsx → kicker", "Termin wählen:"),
    ("termin/page.tsx → Headline Zeile 1", "Ihre 45 Minuten"),
    ("termin/page.tsx → Headline Zeile 2", "Klarheit."),
    (
        "termin/page.tsx → Lead (Teil 1)",
        "Wählen Sie einen Termin, der in Ihren Kalender passt. Danach wissen Sie,",
    ),
    (
        "termin/page.tsx → Lead (hervorgehoben)",
        "wo KI in Ihrem Betrieb entlastet",
    ),
    (
        "termin/page.tsx → Lead (Teil 2)",
        "– und welcher nächste Schritt wirklich sinnvoll ist.",
    ),
    (
        "copy.ts → triggers.terminIntro",
        "Freie Termine · Erstgespräch kostenlos · Keine Vorbereitung nötig",
    ),
    (
        "termin/page.tsx → iframe title",
        "Terminwahl – Erstgespräch Frank Vullhorst",
    ),
    ("termin/page.tsx → Fallback Text", "Kalender lädt nicht?"),
    (
        "termin/page.tsx → Fallback Link",
        "Terminwahl in neuem Tab öffnen",
    ),
    # Impressum
    ("impressum/page.tsx → Zurück-Link", "← Zur Startseite"),
    ("impressum/page.tsx → kicker", "Rechtliches:"),
    ("impressum/page.tsx → Titel", "Impressum"),
    ("impressum/page.tsx → Lead", "Angaben gemäß § 5 TMG"),
    ("impressum/page.tsx → Name", "Frank Vullhorst"),
    ("impressum/page.tsx → Adresse Zeile 1", "Claudiusweg 9"),
    ("impressum/page.tsx → Adresse Zeile 2", "64380 Roßdorf"),
    ("impressum/page.tsx → Adresse Zeile 3", "Deutschland"),
    ("impressum/page.tsx → Kontakt Label", "Kontakt"),
    ("impressum/page.tsx → E-Mail", "Info@FrankVullhorst.de"),
    ("impressum/page.tsx → Telefon", "+49 (0)172 6689960"),
    (
        "impressum/page.tsx → Hinweis",
        "Hinweis: Dieses Impressum ist eine schlanke Basisseite für den Projektstart. Falls weitere rechtliche Angaben benötigt werden, können sie ergänzt werden.",
    ),
    # Datenschutz
    ("datenschutz/page.tsx → Zurück-Link", "← Zur Startseite"),
    ("datenschutz/page.tsx → kicker", "Rechtliches:"),
    ("datenschutz/page.tsx → Titel", "Datenschutz"),
    (
        "datenschutz/page.tsx → Lead",
        "Kurzhinweis zum Datenschutz (DSGVO)",
    ),
    ("datenschutz/page.tsx → Abschnitt 1 Titel", "Verantwortlicher"),
    (
        "datenschutz/page.tsx → Abschnitt 1 Text",
        "Frank Vullhorst / Claudiusweg 9, 64380 Roßdorf / E-Mail: Info@FrankVullhorst.de",
    ),
    (
        "datenschutz/page.tsx → Abschnitt 2 Titel",
        "Terminbuchung über Google",
    ),
    (
        "datenschutz/page.tsx → Abschnitt 2 Text",
        "Die Terminbuchung erfolgt über Google Calendar Appointment Schedules, eingebettet auf der Seite „Termin buchen“ (iframe). Dabei können durch Google personenbezogene Daten verarbeitet werden.",
    ),
    (
        "datenschutz/page.tsx → Hinweis",
        "Hinweis: Diese Datenschutzseite kann bei Bedarf um vollständige Informationen ergänzt werden.",
    ),
]


def main() -> None:
    out = Path(__file__).resolve().parents[1] / "website-texte.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Website-Texte"

    headers = ("Quelle (Datei / Bereich)", "Aktueller Text", "Neuer Text")
    header_fill = PatternFill("solid", fgColor="DDDFEC")
    header_font = Font(bold=True)

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row_idx, (source, current) in enumerate(ROWS, start=2):
        ws.cell(row=row_idx, column=1, value=source).alignment = Alignment(
            vertical="top", wrap_text=True
        )
        ws.cell(row=row_idx, column=2, value=current).alignment = Alignment(
            vertical="top", wrap_text=True
        )
        ws.cell(row=row_idx, column=3, value="").alignment = Alignment(
            vertical="top", wrap_text=True
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{len(ROWS) + 1}"
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 72
    ws.column_dimensions["C"].width = 72

    wb.save(out)
    print(f"Gespeichert: {out} ({len(ROWS)} Zeilen)")


if __name__ == "__main__":
    main()
